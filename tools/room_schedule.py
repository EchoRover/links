#!/usr/bin/env python3
"""Pivot every cohort's schedule by ROOM, which is what Priya asked for.

She sent the workbooks and asked for room-wise schedules in return. The
cohort sheets answer "where is Year 3 CSE on Tuesday"; this answers "what is
in M4-1-017 all week", which is the question you need to book a room, find a
free one, or notice a double-booking.

Two things it does that a hand-made pivot would not:

  * a class taught to several cohorts at once appears ONCE, listing them all.
    Across the current sheets AHUL256, AHUL261 and AGRL130 are each taught to
    Year 3 CSE and EEN together, and Year 1 shares most of its courses across
    all four branches. Listing those separately would read as a clash.
  * it flags REAL clashes: same room, overlapping time, different course.

Writes an .xlsx because that is what the office works in.

    python3 tools/room_schedule.py [out.xlsx]
"""
import collections
import itertools
import re
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from read_xlsx import DAYS, all_cohorts        # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("needs openpyxl:  pip install openpyxl")

ROOT = pathlib.Path(__file__).resolve().parent.parent

# cohort sheet name -> something a human reads
def label(sheet):
    s = sheet.replace(" (2)", "")
    year = {"24": "Y3", "25": "Y2", "26": "Y1"}.get(s[:2], s[:2])
    branch = s[4:7]
    sem = s.split("SEM")[-1] if "SEM" in s else ""
    return f"{year} {branch}{' S' + sem if sem else ''}"


def mins(t):
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def collect():
    rows = []
    for sheet, d in all_cohorts().items():
        for b in d["blocks"]:
            if not b["room"]:
                continue
            rows.append({**b, "cohort": label(sheet)})
    return rows


def merge_shared(rows):
    """One class taught to several cohorts is ONE booking, not several."""
    # NOT keyed on kind. The same class is captioned "Tutorial" on one cohort's
    # sheet and left bare on another's - AHUL261 Monday 09:00 in M4-1-017 is
    # captioned by Y3 CSE and not by Y3 EEN - and keying on kind split one
    # booking into two rows that looked like a room booked against itself.
    by = collections.defaultdict(list)
    for r in rows:
        by[(r["room"], r["day"], r["start"], r["end"], r["course"])].append(r)

    out, disputed = [], []
    for (room, day, start, end, course), group in by.items():
        cohorts = sorted({g["cohort"] + ("" if g["group"] == 0 else f" G{g['group']}")
                          for g in group})
        kinds = {g["kind"] for g in group}
        # a caption beats no caption: if any sheet says Tutorial, it is one
        named = kinds - {"lecture"}
        kind = sorted(named)[0] if named else "lecture"
        if len(named) > 1:
            disputed.append((room, day, start, course, sorted(kinds)))
        out.append({"room": room, "day": day, "start": start, "end": end,
                    "course": course, "kind": kind, "cohorts": cohorts,
                    "kind_disputed": sorted(kinds) if len(kinds) > 1 else None})
    out.sort(key=lambda r: (r["room"], DAYS.index(r["day"]), mins(r["start"])))
    return out


def name_conflicts(rows):
    """A room code and the name printed beside it must agree across sheets.

    Only a cross-sheet view can see this. It has already found one: Y2 EEN
    prints code M4-0-019 with the name "Classroom 8" for AMLL1001, while every
    other sheet and the door call M4-0-019 Classroom 5 and call Classroom 8
    M4-1-011. One of the two is a typo, and which one decides whether that
    17:00 slot collides with the Masters class in the same room.
    """
    def norm(n):
        # "Classroom  6" and "Classroom 6" are the same room typed twice, and
        # so are "BiologyLab" and "Biology Lab". Those are not worth anyone's
        # attention; a room called both Classroom 3 and Classroom 5 is.
        return re.sub(r"\s+", "", n).lower()

    seen = collections.defaultdict(dict)
    for r in rows:
        if r["room"] and r.get("room_name"):
            seen[r["room"]].setdefault(norm(r["room_name"]), r["room_name"])
    return {code: sorted(v.values()) for code, v in seen.items() if len(v) > 1}


def impossible(rows):
    """Blocks whose end is not after their start. A typo, but a silent one."""
    return [r for r in rows if mins(r["end"]) <= mins(r["start"])]


def clashes(bookings):
    bad = []
    for a, b in itertools.combinations(bookings, 2):
        if a["room"] != b["room"] or a["day"] != b["day"] or a["course"] == b["course"]:
            continue
        if mins(a["start"]) < mins(b["end"]) and mins(b["start"]) < mins(a["end"]):
            bad.append((a, b))
    return bad


HEAD = PatternFill("solid", fgColor="1F3864")
ALT = PatternFill("solid", fgColor="EEF2F9")
BAD = PatternFill("solid", fgColor="FFD6D6")


def main():
    out = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else
                       ROOT / "data" / "schedules" / "room-wise-schedule.xlsx")
    rows = collect()
    bookings = merge_shared(rows)
    bad = clashes(bookings)
    naming = name_conflicts(rows)
    broken = impossible(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "By room"
    cols = ["Room", "Day", "Start", "End", "Course", "Type", "Cohorts"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=c).fill = HEAD

    clashed = {id(x) for pair in bad for x in pair}
    room_now, shade = None, False
    for b in bookings:
        if b["room"] != room_now:
            room_now, shade = b["room"], not shade
        ws.append([b["room"], b["day"], b["start"], b["end"], b["course"],
                   "" if b["kind"] == "lecture" else b["kind"],
                   ", ".join(b["cohorts"])])
        if shade or id(b) in clashed:
            fill = BAD if id(b) in clashed else ALT
            for c in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=c).fill = fill

    for i, w in enumerate([12, 11, 8, 8, 11, 9, 34], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Room x Day")
    ws2.append(["Room"] + DAYS)
    for c in range(1, 7):
        ws2.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
        ws2.cell(row=1, column=c).fill = HEAD
    by_room = collections.defaultdict(lambda: collections.defaultdict(list))
    for b in bookings:
        by_room[b["room"]][b["day"]].append(b)
    for room in sorted(by_room):
        line = [room]
        for day in DAYS:
            line.append("\n".join(
                f"{x['start']}-{x['end']} {x['course']}"
                + (f" ({x['kind']})" if x["kind"] != "lecture" else "")
                for x in sorted(by_room[room][day], key=lambda y: mins(y["start"]))))
        ws2.append(line)
        for c in range(1, 7):
            ws2.cell(row=ws2.max_row, column=c).alignment = Alignment(
                vertical="top", wrap_text=True)
    ws2.column_dimensions["A"].width = 12
    for col in "BCDEF":
        ws2.column_dimensions[col].width = 26
    ws2.freeze_panes = "B2"

    wb.save(out)
    rooms = len({b["room"] for b in bookings})
    print(f"{out}")
    print(f"  {len(bookings)} bookings across {rooms} rooms, "
          f"from {len(collect())} cohort entries (shared classes merged)")
    print(f"  blocks whose end time is not after their start: {len(broken)}")
    for r in broken:
        print(f"    {r['cohort']} {r['day']} {r['course']} in {r['room']}: "
              f"\"{r['text'].split(' / ')[-2] if ' / ' in r['text'] else r['text']}\"")
    print(f"  rooms whose CODE and NAME disagree between sheets: {len(naming)}")
    for code, names in naming.items():
        print(f"    {code} is called: " + " / ".join(f'"{n}"' for n in names))
    print(f"  clashes (same room, overlapping, different course): {len(bad)}")
    for a, b in bad:
        print(f"    {a['day']} {a['room']}: {a['course']} {a['start']}-{a['end']} "
              f"({', '.join(a['cohorts'])})  vs  {b['course']} {b['start']}-{b['end']} "
              f"({', '.join(b['cohorts'])})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
